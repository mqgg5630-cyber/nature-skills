from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pathlib import Path

OUT = Path('output')
OUT.mkdir(exist_ok=True)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Palette
BG = RGBColor(247, 244, 236)
INK = RGBColor(27, 44, 38)
MUTED = RGBColor(87, 105, 95)
GREEN = RGBColor(31, 92, 70)
DEEP = RGBColor(20, 58, 47)
TEAL = RGBColor(58, 130, 116)
ORANGE = RGBColor(213, 119, 50)
GOLD = RGBColor(229, 173, 77)
PALE = RGBColor(229, 238, 229)
PALE2 = RGBColor(238, 232, 216)
WHITE = RGBColor(255,255,255)
RED = RGBColor(173, 72, 58)
FONT = 'Aptos'
CN = 'Microsoft YaHei'


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = color


def set_cell_fill(cell, color):
    cell.fill.solid(); cell.fill.fore_color.rgb = color


def add_text(slide, x,y,w,h,text,size=18,color=INK,bold=False,font=CN,align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.06, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin)
    tf.vertical_anchor = valign
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=color
    return box


def add_rich(slide,x,y,w,h,parts,size=18,align=PP_ALIGN.LEFT,margin=0.06):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Inches(margin)
    p=tf.paragraphs[0]; p.alignment=align
    for text,color,bold in parts:
        r=p.add_run(); r.text=text; r.font.name=CN; r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold
    return box


def rect(slide,x,y,w,h,fill,line=None,radius=False):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    sh.line.color.rgb = line if line else fill
    if radius:
        sh.adjustments[0]=0.08
    return sh


def line(slide,x1,y1,x2,y2,color=GREEN,width=2,dash=None):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=color; sh.line.width=Pt(width)
    if dash: sh.line.dash_style = dash
    return sh


def title(slide, kicker, headline, n):
    add_text(slide,0.65,0.32,2.8,0.24,kicker.upper(),10,ORANGE,True,align=PP_ALIGN.LEFT)
    add_text(slide,0.65,0.62,11.8,0.56,headline,26,INK,True)
    line(slide,0.65,1.30,12.68,1.30,PALE2,1)
    add_text(slide,12.15,0.34,0.55,0.22,f'{n:02d}',10,MUTED,True,align=PP_ALIGN.RIGHT)


def footer(slide, text='机器学习筛选鲜味肽｜研究方案汇报'):
    add_text(slide,0.68,7.17,9.5,0.18,text,8,MUTED)
    add_text(slide,11.65,7.17,1.0,0.18,'可编辑方案稿',8,MUTED,align=PP_ALIGN.RIGHT)


def bullet_list(slide,x,y,w,items,size=16,color=INK,gap=0.58,bullet_color=ORANGE):
    for i,it in enumerate(items):
        yy=y+i*gap
        sh=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x),Inches(yy+0.08),Inches(0.11),Inches(0.11)); sh.fill.solid(); sh.fill.fore_color.rgb=bullet_color; sh.line.color.rgb=bullet_color
        add_text(slide,x+0.20,yy,w-0.2,0.42,it,size,color)


def tag(slide,x,y,w,text,fill=PALE,color=GREEN):
    rect(slide,x,y,w,0.30,fill,fill,True); add_text(slide,x+0.06,y+0.03,w-0.12,0.21,text,9,color,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)


def add_chart(slide,x,y,w,h,cats,series,colors,chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED,show_legend=False):
    data=CategoryChartData(); data.categories=cats
    for name,vals in series: data.add_series(name,vals)
    chart=slide.shapes.add_chart(chart_type,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
    chart.has_legend=show_legend
    if show_legend: chart.legend.position=XL_LEGEND_POSITION.BOTTOM; chart.legend.include_in_layout=False
    chart.value_axis.has_major_gridlines=True; chart.value_axis.major_gridlines.format.line.color.rgb=PALE2
    chart.value_axis.tick_labels.font.size=Pt(9); chart.value_axis.tick_labels.font.name=FONT
    chart.category_axis.tick_labels.font.size=Pt(9); chart.category_axis.tick_labels.font.name=CN
    for idx,s in enumerate(chart.series):
        s.format.fill.solid(); s.format.fill.fore_color.rgb=colors[idx%len(colors)]
        s.format.line.color.rgb=colors[idx%len(colors)]
    return chart

# 1 Cover
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s,DEEP)
# molecular motif
for i,(x,y,r,c) in enumerate([(9.9,1.35,.28,GOLD),(11.0,2.0,.20,TEAL),(10.2,3.2,.35,ORANGE),(11.6,3.4,.25,GOLD),(9.4,4.3,.20,TEAL),(11.2,5.0,.32,ORANGE)]):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x),Inches(y),Inches(r),Inches(r)); sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.color.rgb=c
line(s,10.15,1.62,11.08,2.1,GOLD,2); line(s,11.1,2.15,10.42,3.35,TEAL,2); line(s,10.5,3.5,11.7,3.55,ORANGE,2); line(s,10.3,3.55,9.55,4.4,GOLD,2); line(s,9.6,4.5,11.25,5.1,TEAL,2)
add_text(s,0.78,0.72,4,0.28,'NATURE SKILLS · METHODS STORY',11,GOLD,True)
add_text(s,0.78,1.55,8.9,1.25,'机器学习筛选鲜味肽',34,WHITE,True)
add_text(s,0.82,2.98,7.2,0.58,'从序列数据到可验证候选物的完整研究流程',20,RGBColor(215,228,216))
add_text(s,0.82,5.76,6.0,0.35,'汇报人：________    日期：2026 / 09 / 09',12,RGBColor(190,210,197))
add_text(s,0.82,6.65,4,0.22,'研究方案｜方法学汇报｜可编辑 PPTX',9,RGBColor(160,190,172))

# 2 claim
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'01 · 研究问题','鲜味肽筛选的瓶颈，不在于候选太少，而在于“可比、可解释、可验证”',2)
add_text(s,0.78,1.72,7.4,0.88,'目标：建立一个面向短肽序列的、可解释的优先级排序系统，将有限的实验资源投入最有希望的候选物。',23,DEEP,True)
rect(s,0.8,3.05,3.55,2.25,PALE,PALE,True); add_text(s,1.05,3.35,2.9,0.3,'输入',12,GREEN,True); add_text(s,1.05,3.83,2.9,0.7,'肽序列 + 鲜味标签\n实验条件与来源',18,INK,True)
rect(s,4.72,3.05,3.55,2.25,PALE2,PALE2,True); add_text(s,4.97,3.35,2.9,0.3,'学习',12,ORANGE,True); add_text(s,4.97,3.83,2.9,0.7,'特征 → 模型 → 解释\n校准不确定性',18,INK,True)
rect(s,8.64,3.05,3.55,2.25,RGBColor(235,226,218),RGBColor(235,226,218),True); add_text(s,8.89,3.35,2.9,0.3,'输出',12,RED,True); add_text(s,8.89,3.83,2.9,0.7,'候选优先级\n体外验证清单',18,INK,True)
line(s,4.38,4.17,4.70,4.17,ORANGE,3); line(s,8.30,4.17,8.62,4.17,ORANGE,3)
add_text(s,0.82,6.05,11.6,0.36,'核心判断：模型负责缩小搜索空间，实验负责确认真实鲜味。',15,GREEN,True)
footer(s)

# 3 landscape
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'02 · 研究对象','鲜味肽是一个“序列—构象—受体—感知”多尺度问题',3)
# axis
line(s,1.0,4.55,12.0,4.55,GREEN,2)
levels=[('序列',1.15,'氨基酸组成\n长度 / 电荷','输入层',PALE),('理化性质',3.45,'疏水性 / 极性\n分子量 / pI','表征层',PALE2),('受体互作',5.85,'T1R1/T1R3\n结合与构象','机制层',RGBColor(235,226,218)),('感官表型',8.25,'鲜味强度\n协同 / 阈值','输出层',PALE),('应用场景',10.65,'食品基质\n稳定性 / 安全性','转化层',PALE2)]
for name,x,body,sub,c in levels:
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(4.26),Inches(.58),Inches(.58)); sh.fill.solid(); sh.fill.fore_color.rgb=GREEN; sh.line.color.rgb=GREEN
    add_text(s,x-0.2,2.42,1.6,0.35,name,18,INK,True,align=PP_ALIGN.CENTER)
    add_text(s,x-0.2,3.12,1.6,0.72,body,13,MUTED,align=PP_ALIGN.CENTER)
    tag(s,x-0.27,5.00,1.15,sub,c,GREEN)
add_text(s,0.82,1.72,11.3,0.4,'本项目把“鲜味”定义为监督学习标签，但保留实验条件作为可追溯元数据。',16,DEEP,True)
footer(s,'研究对象：短肽序列｜标签、条件、来源必须同时进入数据字典')

# 4 workflow
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'03 · 总体流程','从数据整理到体外验证，形成可回溯的闭环',4)
steps=[('01','数据盘点','来源、标签、条件'),('02','质量控制','去重、冲突、泄漏'),('03','特征工程','序列 + 理化 + 结构'),('04','建模比较','基线、集成、校准'),('05','解释筛选','贡献、适用域、不确定性'),('06','实验验证','合成、感官、机制')]
for i,(num,head,body) in enumerate(steps):
    x=0.82+i*2.03; y=2.15+(i%2)*2.14
    rect(s,x,y,1.62,1.28,PALE if i%2==0 else PALE2,PALE if i%2==0 else PALE2,True)
    add_text(s,x+0.12,y+0.10,0.42,0.25,num,11,ORANGE,True)
    add_text(s,x+0.12,y+0.45,1.36,0.30,head,15,INK,True)
    add_text(s,x+0.12,y+0.83,1.38,0.30,body,10,MUTED)
    if i<5:
        nx=x+1.66; ny=y+0.62
        line(s,nx,ny,nx+0.35,ny,ORANGE,2)
# loop arrow-ish
line(s,11.2,5.60,2.45,5.60,TEAL,2)
add_text(s,0.83,6.18,11.5,0.34,'闭环原则：验证结果回流数据集，更新标签置信度与下一轮主动学习。',16,GREEN,True)
footer(s)

# 5 data
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'04 · 数据工程','先建立“可审计”的数据集，再谈模型性能',5)
# table
rows,cols=6,4; table=s.shapes.add_table(rows,cols,Inches(0.82),Inches(1.75),Inches(7.15),Inches(3.75)).table
widths=[1.35,2.3,1.75,1.75]
for i,w in enumerate(widths): table.columns[i].width=Inches(w)
headers=['字段','内容','处理规则','保留证据']
for c,h in enumerate(headers):
    cell=table.cell(0,c); set_cell_fill(cell,GREEN); cell.text=h
    for p in cell.text_frame.paragraphs: p.runs[0].font.color.rgb=WHITE; p.runs[0].font.bold=True; p.runs[0].font.size=Pt(12); p.runs[0].font.name=CN
vals=[('序列','标准化一字母代码','去空格、统一大小写','原始序列'),('标签','鲜味 / 非鲜味 / 强度','保留原始量表','文献或实验来源'),('条件','pH、温度、基质','缺失不静默填补','条件字典'),('来源','论文、数据库、批次','按来源分组','DOI / accession'),('冲突','同序列异标签','标记冲突，不强行投票','冲突日志')]
for r,row in enumerate(vals,1):
    for c,val in enumerate(row):
        cell=table.cell(r,c); set_cell_fill(cell,WHITE if r%2 else PALE); cell.text=val
        for p in cell.text_frame.paragraphs: p.runs[0].font.color.rgb=INK; p.runs[0].font.size=Pt(10); p.runs[0].font.name=CN
rect(s,8.45,1.78,3.75,1.1,RGBColor(235,226,218),RGBColor(235,226,218),True); add_text(s,8.72,2.00,3.2,0.32,'最大风险：数据泄漏',16,RED,True); add_text(s,8.72,2.43,3.1,0.28,'同源肽 / 同实验批次跨越 train-test',11,INK)
rect(s,8.45,3.17,3.75,1.1,PALE,PALE,True); add_text(s,8.72,3.39,3.2,0.32,'拆分策略：按序列同源簇',16,GREEN,True); add_text(s,8.72,3.82,3.1,0.28,'外层测试集只在最后一次使用',11,INK)
rect(s,8.45,4.56,3.75,1.1,PALE2,PALE2,True); add_text(s,8.72,4.78,3.2,0.32,'输出：数据卡片',16,ORANGE,True); add_text(s,8.72,5.21,3.1,0.28,'版本、来源、标签定义、缺失机制',11,INK)
footer(s,'数据卡片与冲突日志是结果可信度的一部分，不是附属材料')

# 6 features
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'05 · 特征工程','用三类互补表征，避免把鲜味压缩成单一“黑箱分数”',6)
cols=[(0.9,'序列特征','k-mer / one-hot\n预训练蛋白语言模型','可表达局部模式',PALE,GREEN),(4.45,'理化特征','长度、净电荷\n疏水性、pI、分子量','便于解释与审计',PALE2,ORANGE),(8.0,'结构与互作','二级结构倾向\n受体对接 / 接触图','连接机制假设',RGBColor(235,226,218),RED)]
for x,head,body,desc,c,accent in cols:
    rect(s,x,2.0,3.0,3.2,c,c,True); add_text(s,x+0.25,2.32,2.5,0.38,head,18,INK,True); add_text(s,x+0.25,3.1,2.5,0.8,body,17,INK); line(s,x+0.25,4.35,x+2.72,4.35,accent,2); add_text(s,x+0.25,4.62,2.5,0.35,desc,12,MUTED)
add_text(s,0.95,5.95,11.0,0.4,'建模顺序：可解释基线 → 集成模型 → 融合高维表示；每一步都保留消融对照。',16,GREEN,True)
footer(s)

# 7 models
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'06 · 模型比较','性能、稳定性与校准，比单次最高分更重要',7)
add_chart(s,0.9,1.85,6.2,3.8,['逻辑回归','随机森林','XGBoost','ESM+MLP','集成模型'],[('AUROC（示意）',[0.72,0.78,0.81,0.84,0.86])],[TEAL,ORANGE,GREEN,RED,GOLD],XL_CHART_TYPE.BAR_CLUSTERED,False)
add_text(s,1.10,5.78,5.8,0.24,'图示为模型比较的展示方式，数值需由真实数据运行后替换。',9,MUTED,italic=True)
rect(s,7.65,1.92,4.55,3.50,PALE,PALE,True)
add_text(s,7.96,2.24,3.8,0.32,'推荐评估框架',16,GREEN,True)
bullet_list(s,7.96,2.86,3.9,['外层：同源簇留出测试','内层：嵌套交叉验证','指标：AUROC / AUPRC / F1','校准：Brier + reliability','报告：均值 ± 置信区间'],14,gap=0.46)
add_text(s,0.92,6.42,11.4,0.35,'判定规则：若复杂模型仅带来边际收益，优先选择更易解释、部署成本更低的模型。',15,DEEP,True)
footer(s)

# 8 validation
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'07 · 验证设计','把“泛化能力”拆成三个可回答的问题',8)
items=[('能否识别？','区分鲜味与非鲜味\nAUROC / AUPRC',GREEN,PALE),('能否排序？','高分候选是否更富集\nPrecision@K / lift',ORANGE,PALE2),('能否相信？','分数是否可校准\n置信区间 / 适用域',RED,RGBColor(235,226,218))]
for i,(h,b,c,fill) in enumerate(items):
    x=0.9+i*4.08; rect(s,x,2.02,3.45,2.6,fill,fill,True); add_text(s,x+0.3,2.38,2.85,0.40,h,21,c,True); add_text(s,x+0.3,3.18,2.85,0.8,b,17,INK); add_text(s,x+0.3,4.14,2.85,0.25,['分类层','排序层','决策层'][i],10,MUTED,True)
line(s,4.35,5.35,8.85,5.35,ORANGE,2)
add_text(s,0.95,5.67,11.2,0.43,'关键控制：按同源簇拆分，避免“记住相似序列”造成虚高性能。',17,DEEP,True)
footer(s)

# 9 explain
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'08 · 可解释性','解释直接决定候选物是否值得进入实验',9)
# sequence heatmap-like
add_text(s,0.85,1.78,4.5,0.28,'示意：序列贡献图',13,MUTED,True)
seq='G  L  E  P  F  V  D  K  G  R'
for i,ch in enumerate(seq.split('  ')):
    c=[PALE,PALE2,RGBColor(245,211,177),RGBColor(233,184,157),PALE,RGBColor(214,232,220),RGBColor(235,226,218),PALE2,PALE,RGBColor(221,237,228)][i]
    rect(s,0.88+i*0.45,2.38,0.38,0.52,c,c,True); add_text(s,0.88+i*0.45,2.52,0.38,0.18,ch,12,INK,True,align=PP_ALIGN.CENTER)
add_text(s,0.92,3.28,4.4,0.36,'贡献方向与大小应与理化特征、同源序列和实验现象交叉核对。',13,INK)
rect(s,6.0,1.90,6.1,3.35,WHITE,PALE2,True); add_text(s,6.32,2.22,5.2,0.32,'候选入选的四个门槛',17,GREEN,True)
bullet_list(s,6.32,2.86,5.1,['预测分数高：模型排序靠前','不确定性低：落在适用域内','解释一致：特征贡献可复核','实验可行：合成、稳定、安全性可控'],15,gap=0.54)
add_text(s,0.90,5.93,11.5,0.4,'输出应同时包含：分数、置信度、解释与证据链。',17,ORANGE,True)
footer(s)

# 10 screening
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'09 · 虚拟筛选','用“高分 × 低不确定性 × 可实验性”形成候选优先级',10)
# scatter
rect(s,0.92,1.8,6.2,3.95,WHITE,PALE2,True); add_text(s,1.18,2.04,4.5,0.25,'候选排序示意',12,MUTED,True)
line(s,1.45,5.25,6.65,5.25,MUTED,1); line(s,1.45,5.25,1.45,2.55,MUTED,1)
add_text(s,5.55,5.38,1.25,0.2,'预测分数 →',9,MUTED); add_text(s,0.98,2.4,0.35,0.7,'不确定性\n↓',9,MUTED)
pts=[(2.1,4.7,TEAL),(2.8,4.1,GOLD),(3.3,3.8,TEAL),(4.1,4.4,RED),(4.7,3.2,GREEN),(5.2,3.5,GREEN),(5.9,2.9,GREEN),(6.3,2.7,ORANGE)]
for i,(x,y,c) in enumerate(pts):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.18),Inches(.18)); sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.color.rgb=c
    if i in [4,5,6]: add_text(s,x-0.04,y-0.28,0.5,0.18,f'C{i+1}',9,c,True)
rect(s,7.62,1.9,4.55,3.7,PALE,PALE,True); add_text(s,7.92,2.22,3.8,0.30,'建议的候选分层',17,GREEN,True)
for i,(h,b,c) in enumerate([('A类：优先验证','高分、低不确定性、可合成',GREEN),('B类：机制探索','高分但解释或适用域存疑',ORANGE),('C类：暂缓','低分或实验不可行',MUTED)]):
    yy=2.92+i*0.78; add_text(s,7.94,yy,3.6,0.24,h,14,c,True); add_text(s,7.94,yy+0.28,3.55,0.22,b,11,INK)
add_text(s,0.95,6.28,11.6,0.32,'主动学习下一轮：优先补测“模型最不确定、但潜在价值高”的样本。',15,DEEP,True)
footer(s)

# 11 experiment
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'10 · 体外验证','实验验证要同时回答：有没有鲜味、强不强、为什么',11)
phases=[('合成与质控','纯度、分子量\n溶解性、稳定性',PALE,GREEN),('感官 / 细胞','三点检验或强度评分\n受体细胞报告体系',PALE2,ORANGE),('机制补充','受体结合 / 对接\n关键残基突变验证',RGBColor(235,226,218),RED),('迭代更新','回填真实标签\n更新模型与优先级',PALE,GREEN)]
for i,(h,b,fill,c) in enumerate(phases):
    x=0.82+i*3.08; rect(s,x,2.15,2.55,2.3,fill,fill,True); add_text(s,x+0.22,2.48,2.1,0.35,h,17,c,True); add_text(s,x+0.22,3.22,2.1,0.75,b,14,INK)
    if i<3: line(s,x+2.58,3.3,x+3.0,3.3,ORANGE,2)
add_text(s,0.85,5.30,11.5,0.3,'最小验证集建议：覆盖高分、边界、低分对照与模型高不确定性样本。',16,DEEP,True)
add_text(s,0.85,5.93,11.5,0.28,'避免只验证“模型最喜欢的样本”——否则无法区分学习有效与筛选偏差。',13,RED,True)
footer(s)

# 12 risks
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'11 · 风险与边界','模型可以缩小搜索空间，但不能替代鲜味的真实测量',12)
rows,cols=5,3; table=s.shapes.add_table(rows,cols,Inches(0.9),Inches(1.9),Inches(11.4),Inches(3.7)).table
for i,w in enumerate([2.55,4.25,4.6]): table.columns[i].width=Inches(w)
for c,h in enumerate(['风险','可能后果','应对']):
    cell=table.cell(0,c); set_cell_fill(cell,DEEP); cell.text=h
    for p in cell.text_frame.paragraphs: p.runs[0].font.color.rgb=WHITE; p.runs[0].font.bold=True; p.runs[0].font.size=Pt(12); p.runs[0].font.name=CN
risk=[('标签异质性','不同感官量表不可直接合并','分层建模 / 条件编码 / 置信标签'),('同源泄漏','测试性能虚高','同源簇拆分 + 外部留出'),('分布外序列','高分但不可泛化','适用域 + 不确定性筛选'),('基质与协同','纯肽有效，食品中失效','加入基质与组合验证')]
for r,row in enumerate(risk,1):
    for c,val in enumerate(row):
        cell=table.cell(r,c); set_cell_fill(cell,WHITE if r%2 else PALE); cell.text=val
        for p in cell.text_frame.paragraphs: p.runs[0].font.color.rgb=INK; p.runs[0].font.size=Pt(11); p.runs[0].font.name=CN
add_text(s,0.95,6.05,11.4,0.45,'边界声明：本 PPT 是可执行的研究方案；其中模型性能、候选名单与实验结论必须由真实数据产生。',16,RED,True)
footer(s)

# 13 conclusion
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s,DEEP)
add_text(s,0.80,0.70,3,0.25,'12 · TAKE-HOME MESSAGE',11,GOLD,True)
add_text(s,0.80,1.35,11.1,0.95,'把鲜味肽筛选做成\n一个可解释的主动学习闭环。',31,WHITE,True)
line(s,0.82,3.05,12.0,3.05,RGBColor(83,128,103),1)
for i,(num,txt) in enumerate([('01','先管好数据：标签、条件、来源可追溯'),('02','再比较模型：性能、校准、适用域一起看'),('03','最后做实验：验证结果回流，持续更新')]):
    y=3.62+i*0.78; add_text(s,0.85,y,0.45,0.25,num,12,GOLD,True); add_text(s,1.55,y-0.02,9.9,0.34,txt,18,WHITE,True)
add_text(s,0.82,6.60,8.2,0.24,'下一步：锁定数据字典 → 建立基线 → 运行第一轮同源簇验证',11,RGBColor(190,215,196))

# 14 refs
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s,'附录 · 复现与参考','建议把每一轮结果固定为版本化研究记录',14)
add_text(s,0.86,1.70,5.7,0.3,'复现清单',17,GREEN,True)
bullet_list(s,0.86,2.25,5.7,['数据版本与哈希','特征配置与随机种子','拆分索引与外层测试集','模型权重、指标与置信区间','候选清单及实验回填结果'],14,gap=0.52)
add_text(s,7.0,1.70,5.3,0.3,'参考方向（建库时核验）',17,ORANGE,True)
bullet_list(s,7.0,2.25,5.3,['鲜味受体 T1R1/T1R3 与肽类感知研究','机器学习蛋白 / 肽序列表示方法','抗菌肽与生物活性肽预测基准实践','感官评价与食品基质效应方法学'],14,gap=0.52,bullet_color=ORANGE)
add_text(s,0.86,5.75,11.2,0.4,'说明：本方案未虚构数据集规模、模型性能或候选序列；运行真实数据后，应将示意图替换为可追溯结果。',14,RED,True)
footer(s,'汇报完毕｜问题与讨论')

# Add core properties for accessibility and save
prs.core_properties.title='机器学习筛选鲜味肽：从序列到验证的完整流程'
prs.core_properties.subject='Research proposal / methods report'
prs.core_properties.author='Nature Skills'
prs.core_properties.keywords='鲜味肽, 机器学习, 主动学习, 食品科学'
path=OUT/'机器学习筛选鲜味肽_完整流程_可编辑.pptx'
prs.save(path)
print(path)
